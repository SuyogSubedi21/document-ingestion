from pypdf import PdfReader
from docx import Document
from bs4 import BeautifulSoup


def parse_pdf(path):
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def parse_word(path):
    doc = Document(path)
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def parse_html(path):
    with open(path, encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f, "html.parser")
    for tag in soup(["head", "script", "style", "nav", "header", "footer"]):
        tag.decompose()
    return soup.get_text(separator="\n")